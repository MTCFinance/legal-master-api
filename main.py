import os
import re
import secrets
from pathlib import Path
from typing import Optional, List
from uuid import uuid4

from fastapi import FastAPI, Depends, HTTPException, status
from fastapi.responses import FileResponse
from fastapi.security import APIKeyHeader
from pydantic import BaseModel, Field

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from playwright.async_api import (
    async_playwright,
    TimeoutError as PlaywrightTimeoutError,
)


# =========================================================
# FASTAPI APP
# =========================================================

app = FastAPI(
    title="Legal Master API",
    version="1.7.0",
    description="API për Hartues Aktesh Juridike - Kosovë / ARBK",
    servers=[
        {
            "url": "https://legal-master-api.onrender.com",
            "description": "Production server",
        }
    ],
)


# =========================================================
# CONFIG
# =========================================================

BASE_DIR = Path("/tmp/legal-master")
BASE_DIR.mkdir(parents=True, exist_ok=True)

DOWNLOAD_TOKENS = {}

ARBK_SEARCH_URL = "https://arbk.rks-gov.net/TableSearch"


# =========================================================
# API KEY SECURITY
# =========================================================

api_key_header = APIKeyHeader(
    name="X-API-Key",
    auto_error=False,
)


def verify_api_key(
    api_key: Optional[str] = Depends(api_key_header),
):
    expected_key = os.getenv("LEGAL_MASTER_API_KEY")

    if not expected_key:
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="API key is not configured on the server",
        )

    if (
        api_key is None
        or not secrets.compare_digest(api_key, expected_key)
    ):
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Invalid or missing API key",
        )

    return True


# =========================================================
# MODELS
# =========================================================

class HealthResponse(BaseModel):
    status: str
    service: str


class BusinessRequest(BaseModel):
    company_name: str
    legal_form: str = "SHPK"
    address: Optional[str] = None
    capital: Optional[float] = None
    employees: Optional[int] = None


class DocumentRequest(BaseModel):
    document_type: str
    title: str
    body: str
    filename: Optional[str] = None


class ARBKPerson(BaseModel):
    full_name: str
    role: Optional[str] = None


class ARBKBusinessData(BaseModel):
    business_name: str
    trade_name: Optional[str] = None
    nui: str

    legal_form: Optional[str] = None
    business_status: Optional[str] = None
    registration_date: Optional[str] = None

    address: Optional[str] = None
    municipality: Optional[str] = None

    primary_activity: Optional[str] = None
    other_activities: List[str] = Field(default_factory=list)

    owners: List[ARBKPerson] = Field(default_factory=list)
    directors: List[ARBKPerson] = Field(default_factory=list)

    email: Optional[str] = None
    phone: Optional[str] = None

    source_url: Optional[str] = None
    verification_date: Optional[str] = None


# =========================================================
# HELPERS
# =========================================================

def clean_filename(name: str) -> str:
    name = re.sub(r"[^A-Za-z0-9_\- ]", "", name)
    name = name.strip().replace(" ", "_")

    if not name:
        name = "dokument_juridik"

    return name[:100]


def clean_text(value: Optional[str]) -> Optional[str]:
    if not value:
        return None

    value = re.sub(r"\s+", " ", value).strip()

    if not value:
        return None

    return value


def find_first(patterns, text):
    for pattern in patterns:
        match = re.search(
            pattern,
            text,
            flags=re.IGNORECASE | re.MULTILINE,
        )

        if match:
            return clean_text(match.group(1))

    return None


# =========================================================
# WORD DOCUMENT
# =========================================================

def create_word_document(
    data: DocumentRequest,
    output_path: Path,
):
    document = Document()

    normal_style = document.styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(12)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    title_run = title.add_run(data.title.upper())
    title_run.bold = True
    title_run.font.name = "Times New Roman"
    title_run.font.size = Pt(14)

    document.add_paragraph()

    for block in data.body.split("\n"):
        text = block.strip()

        if not text:
            continue

        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        if text.lower().startswith("neni "):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(text)
            run.bold = True
        else:
            run = paragraph.add_run(text)

        run.font.name = "Times New Roman"
        run.font.size = Pt(12)

    document.save(output_path)


# =========================================================
# ARBK DATA EXTRACTION
# =========================================================

def extract_public_arbk_data(
    text: str,
    nui: str,
):
    normalized = re.sub(r"[ \t]+", " ", text)

    business_status = find_first(
        [
            r"\bStatusi\b\s*:?\s*(Aktiv|Pasiv)",
            r"\b(Aktiv|Pasiv)\b",
        ],
        normalized,
    )

    legal_form = find_first(
        [
            r"(Shoqëri\s+me\s+përgjegjësi\s+të\s+kufizuar)",
            r"\b(SH\.?P\.?K\.?)\b",
            r"(Shoqëri\s+Aksionare)",
            r"\b(SH\.?A\.?)\b",
            r"(Biznes\s+Individual)",
        ],
        normalized,
    )

    email_match = re.search(
        r"\b[A-Z0-9._%+\-]+@[A-Z0-9.\-]+\.[A-Z]{2,}\b",
        text,
        re.IGNORECASE,
    )

    email = email_match.group(0) if email_match else None

    phone = find_first(
        [
            (
                r"(?:Telefon|Telefoni|Tel\.?)\s*:?\s*"
                r"(\+?\d[\d\s\-\/]{6,20})"
            )
        ],
        normalized,
    )

    registration_date = find_first(
        [
            (
                r"(?:Themeluar|Regjistruar|Data e regjistrimit)\s*:?\s*"
                r"(\d{1,2}[\/\.\-]\d{1,2}[\/\.\-]\d{4})"
            )
        ],
        normalized,
    )

    address = find_first(
        [
            (
                r"(?:Adresa|Selia)\s*:?\s*"
                r"(.{3,160}?)(?:\n|Komuna|Telefon|Email|$)"
            )
        ],
        text,
    )

    lines = [
        clean_text(line)
        for line in text.splitlines()
        if clean_text(line)
    ]

    business_name = None
    municipality = None

    for index, line in enumerate(lines):
        if nui in line:
            nearby = lines[
                max(0, index - 8):
                min(len(lines), index + 12)
            ]

            for candidate in nearby:
                if not candidate:
                    continue

                lower = candidate.lower()

                if candidate == nui:
                    continue

                if lower in {
                    "emri",
                    "emri tregtar",
                    "nui",
                    "komuna",
                    "lloji biznesit",
                    "njësitë",
                    "statusi",
                }:
                    continue

                if "rezultatet e bizneseve" in lower:
                    continue

                if (
                    "sh.p.k" in lower
                    or "shpk" in lower
                    or "sh.a" in lower
                    or "biznes individual" in lower
                ):
                    business_name = candidate
                    break

            if business_name:
                break

    for index, line in enumerate(lines):
        if nui in line:
            nearby = lines[
                index:
                min(len(lines), index + 12)
            ]

            for candidate in nearby:
                if not candidate or candidate == nui:
                    continue

                lower = candidate.lower()

                if lower in {
                    "aktiv",
                    "pasiv",
                    "sh.p.k.",
                    "shpk",
                }:
                    continue

                if len(candidate) <= 80:
                    municipality = candidate
                    break

            if municipality:
                break

    owners = []
    directors = []

    owner_section = re.search(
        r"PRONAR[ËE]T?(.*?)(?:P[ËE]RFAQ[ËE]SUES|BORDI|$)",
        text,
        re.IGNORECASE | re.DOTALL,
    )

    if owner_section:
        owner_lines = [
            clean_text(x)
            for x in owner_section.group(1).splitlines()
            if clean_text(x)
        ]

        for item in owner_lines[:10]:
            if (
                len(item) >= 3
                and not re.fullmatch(r"\d+", item)
            ):
                owners.append(
                    {
                        "full_name": item,
                        "role": "Pronar",
                    }
                )

    representative_section = re.search(
        r"P[ËE]RFAQ[ËE]SUES(?:IT)?(.*?)(?:BORDI|PRONAR[ËE]T|$)",
        text,
        re.IGNORECASE | re.DOTALL,
    )

    if representative_section:
        rep_lines = [
            clean_text(x)
            for x in representative_section.group(1).splitlines()
            if clean_text(x)
        ]

        for item in rep_lines[:10]:
            if (
                len(item) >= 3
                and not re.fullmatch(r"\d+", item)
            ):
                directors.append(
                    {
                        "full_name": item,
                        "role": "Përfaqësues",
                    }
                )

    return {
        "business_name": business_name,
        "trade_name": None,
        "nui": nui,
        "legal_form": legal_form,
        "business_status": business_status,
        "registration_date": registration_date,
        "address": address,
        "municipality": municipality,
        "primary_activity": None,
        "other_activities": [],
        "owners": owners,
        "directors": directors,
        "email": email,
        "phone": phone,
    }


# =========================================================
# ROOT
# =========================================================

@app.get(
    "/",
    operation_id="root",
)
def root():
    return {
        "service": "Legal Master API",
        "status": "running",
        "version": "1.7.0",
    }


# =========================================================
# HEALTH
# =========================================================

@app.get(
    "/health",
    response_model=HealthResponse,
    operation_id="health",
    dependencies=[Depends(verify_api_key)],
)
def health():
    return {
        "status": "ok",
        "service": "legal-master-api",
    }


# =========================================================
# BUSINESS VALIDATION
# =========================================================

@app.post(
    "/business/validate",
    operation_id="validate_business",
    dependencies=[Depends(verify_api_key)],
)
def validate_business(data: BusinessRequest):
    issues = []

    if not data.company_name.strip():
        issues.append("Mungon emri i biznesit")

    if data.capital is not None and data.capital < 0:
        issues.append("Kapitali nuk mund të jetë negativ")

    if data.employees is not None and data.employees < 0:
        issues.append(
            "Numri i punëtorëve nuk mund të jetë negativ"
        )

    return {
        "valid": len(issues) == 0,
        "issues": issues,
        "data": data,
    }


# =========================================================
# ARBK DEBUG
# =========================================================

@app.get(
    "/arbk/debug-form",
    operation_id="debug_arbk_form",
    dependencies=[Depends(verify_api_key)],
)
async def debug_arbk_form():

    async with async_playwright() as playwright:
        browser = None

        try:
            browser = await playwright.chromium.launch(
                headless=True,
                args=[
                    "--no-sandbox",
                    "--disable-dev-shm-usage",
                ],
            )

            page = await browser.new_page()

            await page.goto(
                ARBK_SEARCH_URL,
                wait_until="domcontentloaded",
                timeout=60000,
            )

            await page.wait_for_timeout(3000)

            nui_field = page.locator("#bsf-NUI")

            return {
                "url": page.url,
                "title": await page.title(),
                "nui_field_found": (
                    await nui_field.count() > 0
                ),
                "nui_field_visible": (
                    await nui_field.is_visible()
                    if await nui_field.count() > 0
                    else False
                ),
                "nui_placeholder": (
                    await nui_field.get_attribute("placeholder")
                    if await nui_field.count() > 0
                    else None
                ),
                "search_button_count": await page.get_by_role(
                    "button",
                    name="KËRKO",
                    exact=True,
                ).count(),
            }

        finally:
            if browser:
                await browser.close()


# =========================================================
# ARBK PUBLIC LOOKUP
# =========================================================

@app.get(
    "/arbk/lookup/{nui}",
    operation_id="lookup_arbk_business",
    dependencies=[Depends(verify_api_key)],
)
async def lookup_arbk_business(nui: str):

    nui = re.sub(r"\D", "", nui)

    if not nui:
        raise HTTPException(
            status_code=400,
            detail="Numri unik identifikues nuk është valid",
        )

    async with async_playwright() as playwright:
        browser = None

        try:
            browser = await playwright.chromium.launch(
                headless=True,
                args=[
                    "--no-sandbox",
                    "--disable-dev-shm-usage",
                ],
            )

            page = await browser.new_page(
                viewport={
                    "width": 1440,
                    "height": 1200,
                }
            )

            await page.goto(
                ARBK_SEARCH_URL,
                wait_until="domcontentloaded",
                timeout=60000,
            )

            await page.wait_for_timeout(3000)

            # =================================================
            # EXACT NUI FIELD
            # =================================================

            nui_field = page.locator("#bsf-NUI")

            if await nui_field.count() == 0:
                raise HTTPException(
                    status_code=502,
                    detail=(
                        "Fusha #bsf-NUI nuk u gjet "
                        "në faqen e ARBK-së."
                    ),
                )

            if not await nui_field.is_visible():
                raise HTTPException(
                    status_code=502,
                    detail=(
                        "Fusha #bsf-NUI ekziston, "
                        "por nuk është e dukshme."
                    ),
                )

            await nui_field.click()

            await nui_field.fill("")

            await nui_field.fill(nui)

            # Verify that browser really contains the number
            entered_value = await nui_field.input_value()

            if entered_value != nui:
                raise HTTPException(
                    status_code=502,
                    detail=(
                        "Numri nuk u vendos saktë "
                        "në fushën #bsf-NUI."
                    ),
                )

            # =================================================
            # EXACT SEARCH BUTTON
            # =================================================

            search_button = page.get_by_role(
                "button",
                name="KËRKO",
                exact=True,
            )

            if await search_button.count() == 0:
                raise HTTPException(
                    status_code=502,
                    detail=(
                        "Butoni KËRKO nuk u gjet."
                    ),
                )

            await search_button.click()

            # Wait for page/search result
            try:
                await page.wait_for_load_state(
                    "networkidle",
                    timeout=20000,
                )
            except PlaywrightTimeoutError:
                pass

            # Wait until result section appears if possible
            try:
                await page.get_by_text(
                    "Rezultatet e Bizneseve",
                    exact=False,
                ).wait_for(
                    state="visible",
                    timeout=15000,
                )
            except Exception:
                pass

            await page.wait_for_timeout(4000)

            body_text = await page.locator(
                "body"
            ).inner_text()

            # =================================================
            # CHECK RESULT
            # =================================================

            if nui not in body_text:
                return {
                    "found": False,
                    "nui": nui,
                    "source": ARBK_SEARCH_URL,
                    "entered_value": entered_value,
                    "message": (
                        "Kërkimi u ekzekutua në #bsf-NUI, "
                        "por ky numër nuk u gjet në tekstin "
                        "publik të rezultatit."
                    ),
                    "raw_public_text": body_text[:15000],
                }

            extracted = extract_public_arbk_data(
                body_text,
                nui,
            )

            return {
                "found": True,
                "nui": nui,
                "source": ARBK_SEARCH_URL,
                "entered_value": entered_value,
                "data": extracted,
                "raw_public_text": body_text[:15000],
                "note": (
                    "Kërkimi është bërë direkt në fushën "
                    "#bsf-NUI dhe me butonin zyrtar KËRKO. "
                    "Janë përdorur vetëm të dhënat publike."
                ),
            }

        except HTTPException:
            raise

        except PlaywrightTimeoutError:
            raise HTTPException(
                status_code=504,
                detail=(
                    "Faqja e ARBK-së nuk u përgjigj "
                    "brenda afatit të lejuar."
                ),
            )

        except Exception as exc:
            raise HTTPException(
                status_code=502,
                detail=(
                    "Gabim gjatë kërkimit publik në ARBK: "
                    f"{str(exc)}"
                ),
            )

        finally:
            if browser:
                await browser.close()


# =========================================================
# MANUAL ARBK VERIFICATION
# =========================================================

@app.post(
    "/arbk/verify-business",
    operation_id="verify_arbk_business",
    dependencies=[Depends(verify_api_key)],
)
def verify_arbk_business(data: ARBKBusinessData):

    critical_issues = []
    warnings = []

    if not data.business_name.strip():
        critical_issues.append(
            "Mungon emri i saktë i biznesit"
        )

    if not data.nui.strip():
        critical_issues.append(
            "Mungon NUI"
        )

    if data.business_status is None:
        warnings.append(
            "Statusi i biznesit nuk është dhënë"
        )

    if data.address is None:
        warnings.append(
            "Adresa e biznesit nuk është dhënë"
        )

    if data.legal_form is None:
        warnings.append(
            "Forma juridike nuk është dhënë"
        )

    if not data.directors:
        warnings.append(
            "Nuk është dhënë drejtori/përfaqësuesi"
        )

    if not data.owners:
        warnings.append(
            "Nuk janë dhënë pronarët/anëtarët"
        )

    return {
        "valid": len(critical_issues) == 0,
        "critical_issues": critical_issues,
        "warnings": warnings,
        "business": data.model_dump(),
    }


# =========================================================
# DOCUMENT GENERATION
# =========================================================

@app.post(
    "/documents/generate",
    operation_id="generate_document",
    dependencies=[Depends(verify_api_key)],
)
def generate_document(data: DocumentRequest):

    if not data.title.strip():
        raise HTTPException(
            status_code=400,
            detail="Titulli i dokumentit mungon",
        )

    if not data.body.strip():
        raise HTTPException(
            status_code=400,
            detail="Përmbajtja e dokumentit mungon",
        )

    if data.filename:
        base_filename = clean_filename(
            data.filename
        )
    else:
        base_filename = clean_filename(
            f"{data.document_type}_{uuid4().hex[:8]}"
        )

    filename = f"{base_filename}.docx"

    internal_name = (
        f"{uuid4().hex}_{filename}"
    )

    output_path = (
        BASE_DIR / internal_name
    )

    create_word_document(
        data=data,
        output_path=output_path,
    )

    download_token = secrets.token_urlsafe(32)

    DOWNLOAD_TOKENS[download_token] = {
        "path": str(output_path),
        "filename": filename,
    }

    download_url = (
        "https://legal-master-api.onrender.com"
        f"/files/{download_token}"
    )

    return {
        "success": True,
        "document_type": data.document_type,
        "filename": filename,
        "download_url": download_url,
    }


# =========================================================
# FILE DOWNLOAD
# =========================================================

@app.get(
    "/files/{token}",
    operation_id="download_document",
)
def download_document(token: str):

    file_data = DOWNLOAD_TOKENS.get(token)

    if not file_data:
        raise HTTPException(
            status_code=404,
            detail=(
                "Dokumenti nuk ekziston "
                "ose linku ka skaduar"
            ),
        )

    file_path = Path(
        file_data["path"]
    )

    if not file_path.exists():
        raise HTTPException(
            status_code=404,
            detail=(
                "Dokumenti nuk gjendet "
                "më në server"
            ),
        )

    return FileResponse(
        path=file_path,
        filename=file_data["filename"],
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        ),
    )
